"""Source-defined catalogue-entry accounting; unknown records are never scaled.

Outputs distinguish raw inputs, selected entries, mapped ground, space, unknown,
and documented exclusions. Selected entries are NOT certified unique exposures.
"""
from pathlib import Path
from collections import Counter, defaultdict
from decimal import Decimal
import csv
import hashlib
import html
import json
import math
import os
import re
import openpyxl

from paths import ROOT, INPUTS, RESOURCES, DATA
CATEGORIES = ['J6-13','S1-8','S9','U1-5','N1','N2']
NCOLS = dict(zip(CATEGORIES,[2,8,1,5,1,1]))
POLICY = json.loads((RESOURCES/'counting_policy.json').read_text(encoding='utf-8'))
RAW_DEFAULT = r'C:\Users\Y\Documents\study_ancoj\data\imcce' if os.name=='nt' else '/mnt/c/Users/Y/Documents/study_ancoj/data/imcce'


def read_text(path):
    data = Path(path).read_bytes()
    try: return data.decode('utf-8-sig')
    except UnicodeDecodeError: return data.decode('cp1252')


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def normalize_code(value):
    s=str(value).strip().upper()
    if s.isdigit() and len(s)<=3: s=s.zfill(3)
    return s if re.fullmatch(r'[A-Z0-9]{3}',s) else None


def geodetic_latitude(rho_cos, rho_sin):
    """Bowring ECEF inversion: retains station height encoded in MPC vector."""
    if rho_cos==0 and rho_sin==0: return None
    a=6378137.0; f=1/298.257223563; b=a*(1-f); e2=f*(2-f)
    p=a*rho_cos; z=a*rho_sin
    theta=math.atan2(z*a,p*b)
    return math.degrees(math.atan2(z+(a*a-b*b)/b*math.sin(theta)**3,
                                   p-e2*a*math.cos(theta)**3))


def load_stations():
    stations={}
    for l in read_text(RESOURCES/'obscodes_snapshot.dat').splitlines():
        if len(l)<30: continue
        try:
            code=l[:3];lon=float(l[4:13]);c=float(l[13:22]);s=float(l[22:32])
        except ValueError: continue
        stations[code]=dict(code=code,name=l[33:].strip(),longitude=((lon+180)%360)-180,
                            latitude=geodetic_latitude(c,s),rho_cos=c,rho_sin=s,
                            coordinate_source='resources/obscodes_snapshot.dat',country='Unknown',country_source='unassigned')
    supplement=json.loads((RESOURCES/'mpc_supplement.json').read_text(encoding='utf-8'))
    for r in supplement['stations']:
        stations[r['code']]=dict(code=r['code'],name=r['name'],longitude=((r['longitude']+180)%360)-180,
                                 latitude=geodetic_latitude(r['rho_cos'],r['rho_sin']),rho_cos=r['rho_cos'],rho_sin=r['rho_sin'],
                                 country=r['country'],coordinate_source=supplement['source_url'],country_source='verified_site_identity')
    country_codes={
        'China':['286','337','327','330','334','324','305'],
        'USA':['786','787','780','754','688','689','690','673','675','644','662','704','702','701','568','608','711','691','695','672','801'],
        'Brazil':['874','883'], 'UK':['000','001','002','501'],
        'France':['003','999','511','586','005','007','020'], 'Germany':['536'],
        'Spain':['950','954','G40'], 'Australia':['413','415','E10','E12'],
        'Ukraine':['083','086','089'], 'Russia':['084','114','115','136','B18'],
        'Uzbekistan':['186','188','192'], 'Georgia':['119'], 'Chile':['807','809','262'],
        'Japan':['381','387','388'], 'Belgium':['012'], 'Romania':['073']}
    for country,codes in country_codes.items():
        for code in codes:
            if code in stations and stations[code]['country']=='Unknown':
                stations[code].update(country=country,country_source='catalogue_site_identity')
    from matplotlib.path import Path as PolygonPath
    features=json.loads((RESOURCES/'ne_110m_admin_0_countries.geojson').read_text(encoding='utf-8'))['features']
    polygons=[]
    for f in features:
        geom=f['geometry'];coords=[geom['coordinates']] if geom['type']=='Polygon' else geom['coordinates']
        polygons.append((f['properties']['ADMIN'],[(PolygonPath(p[0]),[PolygonPath(h) for h in p[1:]]) for p in coords]))
    for st in stations.values():
        if st['country']!='Unknown' or st['latitude'] is None: continue
        pt=(st['longitude'],st['latitude'])
        matches=[country for country,polys in polygons if any(outer.contains_point(pt) and not any(h.contains_point(pt) for h in holes) for outer,holes in polys)]
        if len(matches)==1:
            st.update(country={'United States of America':'USA','United Kingdom':'UK'}.get(matches[0],matches[0]),
                      country_source='Natural Earth 110m polygon containment (inferred)')
    return stations


ALIASES={'la palma':'950','haute provence':'511','pic du midi':'586','sheshan':'337','sheshan station':'337',
         'yunnan':'286','yunnan obs.':'286','yunnan observatory':'286','mcdonald':'711','mccormick':'780',
         'itajuba':'874','lab.nac.astr':'874','lijiang':'O44','lijiang stat.':'O44','greenwich':'000',
         'meudon':'005','pulkovo':'084','tokyo-mitaka (388)':'388','bucharest':'073','uccle':'012',
         'bordeaux-floirac':'999','bordeaux':'999','arequipa':'805','cordoba':'822','siding spring':'413',
         'siding spring observatory':'413','table mountain':'673'}


def metadata(path):
    if not path: return dict(text='',site='',codes=[],station_field=None,format_evidence='')
    t=html.unescape(re.sub(r'<[^>]*>','',read_text(path)))
    m=re.search(r'(?im)^\s*observator(?:y|ies)\s*:\s*([^\n]*(?:\n[ \t]+[^\n]+)*)',t)
    site=m[1].strip() if m else ''
    site=re.split(r'(?im)^\s*(Reference|Informations?|Comments|Format)[.:\s]',site)[0].strip()
    codes=[]
    for l in site.splitlines():
        match=re.match(r'\s*([A-Za-z0-9]{1,3})\s*[-–]\s*',l)
        if match:
            c=normalize_code(match[1])
            if c and c not in codes: codes.append(c)
    field=None;evidence=''
    form=re.split(r'(?im)^\s*Format[.:]',t)
    if len(form)>1:
        for l in form[-1].splitlines():
            m=re.match(r'\s*(\d+)\.\s*(.*)',l)
            if m and (re.search(r'observator',m[2],re.I) or re.search(r'(site.*code|code.*site)',m[2],re.I)):
                if 'Column' not in m[2] and re.search(r'code|number|\bN\b',m[2],re.I):
                    field=int(m[1])-1;evidence=l.strip();break
    return dict(text=t,site=site,codes=codes,station_field=field,format_evidence=evidence)


def read_sources():
    raw=Path(os.environ.get('IMCCE_RAW_DIR',RAW_DEFAULT))
    if not raw.is_dir(): raise FileNotFoundError(f'Raw inputs missing: {raw}; set IMCCE_RAW_DIR.')
    index={p.stem.lower():p for p in raw.rglob('*.txt')}
    index.update({p.stem.lower():p for p in (INPUTS/'supplementary_catalogues').glob('*.txt')})
    htmls={p.stem.lower():p for p in raw.rglob('*.html')}
    wb=openpyxl.load_workbook(INPUTS/'团队观测数据占比.xlsx',data_only=True)
    sources=[]
    for ws in wb:
        if ws.title not in CATEGORIES: continue
        headers=[str(c.value or '').lower() for c in ws[1]]
        fi=next(i for i,h in enumerate(headers) if 'ident' in h)
        si=next(i for i,h in enumerate(headers) if 'site' in h)
        ri=next(i for i,h in enumerate(headers) if 'ref' in h)
        for rn,vals in enumerate(ws.iter_rows(min_row=2,values_only=True),2):
            if not any(vals): continue
            count=sum(int(v) for v in vals[1:NCOLS[ws.title]+1] if isinstance(v,(int,float)))
            if count==0: continue
            fid=str(vals[fi] or '').lower().strip()
            if not fid:
                if ws.title=='U1-5' and count==10339: fid='supplement_zhang2024'
                else: raise ValueError(f'Missing source ID: {ws.title}:{rn}')
            sources.append(dict(sheet=ws.title,file_id=fid,workbook_row=rn,count=count,site_raw=str(vals[si] or ''),
                                component_counts=[int(v) if isinstance(v,(int,float)) else 0 for v in vals[1:NCOLS[ws.title]+1]],
                                reference=POLICY['team_sources'].get(fid,str(vals[ri] or ''))))
    if not any(s['file_id']=='jo0080' for s in sources):
        sources.append(dict(sheet='J6-13',file_id='jo0080',workbook_row=0,count=2354,site_raw='286, O44',reference=POLICY['team_sources']['jo0080']))
    ids=[s['file_id'] for s in sources]
    if len(ids)!=len(set(ids)): raise ValueError('Duplicate source ID in workbook/append list')
    for s in sources:
        fid=s['file_id'];s['raw_path']=str(index.get(fid,''));s['content_path']=str(htmls.get(fid,''));s['meta']=metadata(htmls.get(fid))
        s['lines']=[(n,l) for n,l in enumerate(read_text(index[fid]).splitlines(),1) if l.strip() and not l.lstrip().startswith('#')] if fid in index else []
    return sources


def space_source(source):
    s=source['meta']['site']+' '+source['site_raw']
    return bool(re.search(r'(?i)\b(Gaia|Cassini|Voyager|Voager|Hipparcos)\b|\bC51\b|WISE\s*\(sat|satellite\s+WISE',s))


def source_single_station(source):
    codes=source['meta']['codes']
    if len(codes)==1: return codes[0]
    if len(codes)>1: return None
    for site in [source['meta']['site'],source['site_raw']]:
        if '\n' in site or ',' in site or ';' in site: continue
        code=normalize_code(site)
        if code: return code
        m=re.fullmatch(r'\s*([A-Za-z0-9]{1,3})\s*[-–]\s*[^,;]+',site)
        if m: return normalize_code(m[1])
        if site.lower().strip() in ALIASES: return ALIASES[site.lower().strip()]
    return None


def parse_station(source,line):
    fid=source['file_id'];meta=source['meta'];parts=line.split()
    if fid in ['sm0034','sm0035','sm0036']: return normalize_code(line[27:30]),'fixed_columns_28_30'
    if fid=='jo0080': return normalize_code(parts[-1]),'published_source_last_field'
    if fid=='um0034': return {'P':'586','L':'511','E':'809','H':'568'}.get(parts[0][0]),'source_plate_prefix'
    if fid=='sm0084':
        return {'1':'084','2':'084','3':'119'}.get(parts[4]),'source_telescope_to_site_definition'
    if fid in ['um0062','um0070']:
        # Both source headers distinguish USNO before/since 1893; year is field 2.
        return ('787' if int(parts[1])<1893 else '786'),'source_USNO_historical_site_epoch'
    if meta['station_field'] is not None:
        i=meta['station_field'];code=normalize_code(parts[i]) if i<len(parts) else None
        if meta['codes'] and code not in meta['codes']: return None,'station_not_in_source_site_list'
        return code,'source_format_field_'+str(i+1)
    return source_single_station(source),'source_single_station_header'


def epoch_key(line,schema,station):
    t=line.split()
    try:
        if schema=='ymd_station': return (int(t[0]),int(t[1]),str(Decimal(t[2]).normalize()),station)
        if schema=='jd_satellite': return (str(Decimal(t[0]).normalize()),int(t[1]),station)
    except (ValueError,IndexError,ArithmeticError): return None
    return None


def build(output_dir=DATA):
    output_dir=Path(output_dir);output_dir.mkdir(parents=True,exist_ok=True)
    stations=load_stations();sources=read_sources();by_id={s['file_id']:s for s in sources}
    duplicate_lines=defaultdict(dict)
    for rule in POLICY['exact_epoch_replacements']:
        new_keys={epoch_key(l,rule['schema'],rule['station']) for _,l in by_id[rule['new']]['lines']};new_keys.discard(None)
        for n,l in by_id[rule['old']]['lines']:
            k=epoch_key(l,rule['schema'],rule['station'])
            if k is not None and k in new_keys: duplicate_lines[rule['old']][n]=rule['new']
    ledger=[];unresolved=[];excluded=[];record_rows=[];raw_scope=[];stat=defaultdict(Counter)
    for s in sources:
        fid=s['file_id'];cat=s['sheet'];n=s['count'];team=fid in POLICY['team_sources']
        base={k:s[k] for k in ['sheet','file_id','workbook_row','site_raw','reference','raw_path','content_path']}
        totals=Counter(input=n)
        def record(amount,code,status,reason,line_no=0):
            if not amount: return
            row=dict(sheet=cat,file_id=fid,line=line_no,code=code or '',count=amount,status=status,reason=reason)
            record_rows.append(row);totals[status]+=amount
            if status=='mapped': stat[code][cat]+=amount
            elif status in ['space','unresolved']: unresolved.append(dict(**base,line=line_no,code=code or '',count=amount,status=status,reason=reason))
            else: excluded.append(dict(**base,line=line_no,code=code or '',count=amount,status=status,reason=reason))
        rule=POLICY['full_exclusions'].get(fid)
        if rule: record(n,None,'excluded',rule)
        elif space_source(s): record(n,None,'space','Source identifies space mission; no fixed ground position')
        else:
            lines=s['lines'];field=s['meta']['station_field'];single=source_single_station(s)
            if fid=='um0034':
                # The index has U1/U2/U3 counts, U4 is reference-only, U5 blank.
                # Preserve extra raw Miranda rows separately; never hide the six-row gap.
                targets={i+1 for i,c in enumerate(s['component_counts']) if c>0}
                for ln,l in lines:
                    if int(l.split()[4]) not in targets:
                        raw_scope.append(dict(file_id=fid,line=ln,raw_text=l,
                                              reason='Raw target absent from nonzero catalogue count columns; not included in catalogue-entry denominator'))
                lines=[(ln,l) for ln,l in lines if int(l.split()[4]) in targets]
            if not lines and fid=='supplement_zhang2024':
                record(n,'286','mapped','Published station 286 aggregate; local raw rows unavailable')
            elif single and field is None and fid not in duplicate_lines and fid not in ['sm0034','sm0035','sm0036']:
                st=stations.get(single)
                if st and st['latitude'] is not None: record(n,single,'mapped','Single-station source header; catalogue position-entry aggregate')
                else: record(n,single,'unresolved','Station lacks ground coordinates or has zero vector')
            elif not lines: record(n,None,'unresolved','Missing raw file; no unambiguous single-station source')
            else:
                weight=1
                if len(lines)*2==n and re.search(r'PHESAT|mutual|phen\.mut',s['meta']['text']+' '+s['site_raw'],re.I): weight=2
                if len(lines)*weight>n:
                    record(n,None,'unresolved',f'Raw lines {len(lines)} exceed catalogue count {n}; review unit or target filter')
                else:
                    for ln,line in lines:
                        if ln in duplicate_lines[fid]:
                            record(weight,None,'excluded','known_overlap_exact_epoch: newer '+duplicate_lines[fid][ln],ln);continue
                        code,method=parse_station(s,line);st=stations.get(code)
                        if st and st['latitude'] is not None: record(weight,code,'mapped',method,ln)
                        else: record(weight,code,'unresolved',method+': station or coordinates unverified',ln)
                    record(n-len(lines)*weight,None,'unresolved',f'Catalogue/raw count gap: catalogue={n}, raw={len(lines)}, weight={weight}')
        selected=totals['mapped']+totals['space']+totals['unresolved'];residual=n-selected-totals['excluded']
        if residual: raise AssertionError((fid,residual))
        ledger.append(dict(**base,team=team,input_count=n,selected_count=selected,mapped_count=totals['mapped'],space_count=totals['space'],
                           unresolved_count=totals['unresolved'],excluded_count=totals['excluded'],residual=residual,
                           raw_line_count=len(s['lines']),station_field=s['meta']['station_field'],format_evidence=s['meta']['format_evidence'],source_site=s['meta']['site']))
    station_rows=[]
    for code,c in stat.items():
        st=stations[code]
        station_rows.append(dict(code=code,name=st['name'],country=st['country'],is_china=st['country']=='China',longitude=round(st['longitude'],6),latitude=round(st['latitude'],6),
                                 **{k:c[k] for k in CATEGORIES},total_obs=sum(c.values()),coordinate_source=st['coordinate_source'],country_source=st['country_source']))
    station_rows.sort(key=lambda r:(-r['total_obs'],r['code']))
    countries=defaultdict(Counter)
    for r in station_rows:
        for k in CATEGORIES+['total_obs']:countries[r['country']][k]+=r[k]
        countries[r['country']]['station_count']+=1
    country_rows=[dict(country=k,**dict(v)) for k,v in sorted(countries.items(),key=lambda kv:(-kv[1]['total_obs'],kv[0]))]
    category_rows=[]
    for cat in CATEGORIES:
        subset=[r for r in ledger if r['sheet']==cat];inp=sum(r['input_count'] for r in subset);sel=sum(r['selected_count'] for r in subset)
        tm=sum(r['input_count'] for r in subset if r['team']);ts=sum(r['selected_count'] for r in subset if r['team'])
        category_rows.append(dict(system=cat,raw_input_entries=inp,satellite_entries=inp-(by_id['nm0090']['count'] if cat=='N1' else 0),raw_team_entries=tm,
                                  selected_entries=sel,selected_team_entries=ts,selected_team_pct=100*ts/sel if sel else 0,
                                  mapped_entries=sum(r['mapped_count'] for r in subset),space_entries=sum(r['space_count'] for r in subset),
                                  unresolved_entries=sum(r['unresolved_count'] for r in subset),excluded_entries=sum(r['excluded_count'] for r in subset),residual=sum(r['residual'] for r in subset)))
    def write_csv(name,rows):
        if not rows: raise ValueError('Unexpected empty output: '+name)
        with (output_dir/name).open('w',newline='',encoding='utf-8-sig') as f:
            w=csv.DictWriter(f,fieldnames=list(rows[0]));w.writeheader();w.writerows(rows)
    for name,rows in [('catalogue_source_ledger.csv',ledger),('record_classification.csv',record_rows),('observatory_satellite_summary.csv',station_rows),
                      ('country_satellite_summary.csv',country_rows),('unresolved_records_flagged.csv',unresolved),('excluded_records.csv',excluded),('satellite_statistics.csv',category_rows)]: write_csv(name,rows)
    write_csv('raw_records_outside_catalogue_scope.csv',raw_scope)
    inputs={INPUTS/'团队观测数据占比.xlsx',*RESOURCES.iterdir()}
    for s in sources:
        for key in ['raw_path','content_path']:
            if s[key]: inputs.add(Path(s[key]))
    keys=['raw_input_entries','satellite_entries','raw_team_entries','selected_entries','selected_team_entries','mapped_entries','space_entries','unresolved_entries','excluded_entries','residual']
    manifest=dict(policy=POLICY,source_count=len(sources),input_hashes=[dict(path=str(p),sha256=sha(p)) for p in sorted(inputs) if p.is_file()],
                  totals={k:sum(r[k] for r in category_rows) for k in keys},limitations=POLICY['unresolved_overlap_families']+
                  ['Zhang 2024 aggregate 10339 included; no local raw rows.','Selected entries are not certified unique observations.','Country polygon attribution is labelled as inference.'])
    (output_dir/'pipeline_manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
    return manifest
